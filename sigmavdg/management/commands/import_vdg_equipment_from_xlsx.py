import unicodedata
from datetime import datetime

from django.core.management.base import BaseCommand, CommandError

from openpyxl import load_workbook

from sigmavdg.models import VDGEquipment


def _normalize_header(value):
    if value is None:
        return ""
    text = str(value).strip().lower()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return " ".join(text.split())


def _as_text(value):
    if value is None:
        return ""
    return str(value).strip()


def _parse_date(value):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, str):
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
            try:
                return datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                continue
    return None


def _parse_bool(value):
    if value is None:
        return False
    if isinstance(value, bool):
        return value
    text = str(value).strip().lower()
    return text in {"si", "sí", "s", "yes", "y", "1", "true", "x"}


class Command(BaseCommand):
    help = "Importa equipos VDG desde Excel (sheet 'LISTADO EQUIPOS')."

    def add_arguments(self, parser):
        parser.add_argument("--file", required=True, help="Ruta al Excel de equipos VDG")

    def handle(self, *args, **options):
        file_path = options["file"]
        try:
            workbook = load_workbook(filename=file_path, data_only=True)
        except FileNotFoundError as exc:
            raise CommandError(f"No se encuentra el archivo: {file_path}") from exc
        except Exception as exc:
            raise CommandError(f"No se pudo abrir el archivo: {exc}") from exc

        if "LISTADO EQUIPOS" not in workbook.sheetnames:
            raise CommandError("No se encuentra la hoja 'LISTADO EQUIPOS'.")

        sheet = workbook["LISTADO EQUIPOS"]
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            self.stdout.write(self.style.WARNING("La hoja esta vacia."))
            return

        headers = [_normalize_header(value) for value in rows[0]]
        header_map = {header: idx for idx, header in enumerate(headers) if header}

        def get_value(row, *keys):
            for key in keys:
                idx = header_map.get(key)
                if idx is not None and idx < len(row):
                    return row[idx]
            return None

        imported = 0
        updated = 0
        for row in rows[1:]:
            code = get_value(row, "codigo", "codigo equipo", "cod")
            if not code:
                continue
            code = _as_text(code)
            defaults = {
                "description": _as_text(
                    get_value(row, "descripcion", "descripcion equipo")
                ),
                "is_reference": _parse_bool(
                    get_value(row, "es patron", "patron", "es patron?", "es patron ?")
                ),
                "responsible": _as_text(get_value(row, "responsable")),
                "location": _as_text(get_value(row, "localizacion", "ubicacion")),
                "received_date": _parse_date(get_value(row, "fecha recepcion")),
                "decommission_date": _parse_date(get_value(row, "fecha baja")),
                "observations": _as_text(get_value(row, "observaciones")),
                "brand": _as_text(get_value(row, "marca")),
                "model": _as_text(get_value(row, "modelo")),
                "serial_number": _as_text(get_value(row, "serie", "numero serie")),
                "range": _as_text(get_value(row, "rango")),
                "resolution": _as_text(get_value(row, "resolucion")),
                "tolerance": _as_text(get_value(row, "tolerancia")),
            }

            obj, created = VDGEquipment.objects.update_or_create(
                code=code, defaults=defaults
            )
            if created:
                imported += 1
            else:
                updated += 1

        self.stdout.write(
            self.style.SUCCESS(f"Importados: {imported}, actualizados: {updated}.")
        )
