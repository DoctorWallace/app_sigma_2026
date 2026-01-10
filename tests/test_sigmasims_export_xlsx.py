import pytest
from datetime import date
from io import BytesIO

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.urls import reverse
from openpyxl import load_workbook

from icts.sigmasims.models import SIMSRecord


def _create_user(username, groups):
    User = get_user_model()
    user = User.objects.create_user(
        username=username,
        password="safe-pass",
        email=f"{username}@example.com",
    )
    for name in groups:
        group, _ = Group.objects.get_or_create(name=name)
        user.groups.add(group)
    return user


@pytest.mark.django_db
def test_sigmasims_export_xlsx(client):
    user = _create_user("sims_export", ["tecnicos_sims"])
    record = SIMSRecord.objects.create(
        reception_date=date(2025, 2, 1),
        analysis_date=date(2025, 2, 3),
        sample_identification="Sample X",
        client_name="Client X",
        sample_characteristics="Chars",
        responsible_name="Resp",
        client_requirements="Req",
        incidents="Incidents",
        comments="Comments",
    )

    client.force_login(user)
    response = client.get(reverse("icts:sigmasims:record_export_xlsx"))
    assert response.status_code == 200
    assert response["Content-Type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    workbook = load_workbook(BytesIO(response.content))
    assert "SIMS" in workbook.sheetnames
    sheet = workbook["SIMS"]

    assert (sheet["B9"].value or "").strip()
    assert (sheet["P9"].value or "").strip()
    assert sheet["B11"].value == record.sims_id
    assert sheet["E11"].value == "Sample X"

    formula = sheet["N11"].value
    assert isinstance(formula, str)
    assert "IF(OR(" in formula or "J11-D11" in formula
