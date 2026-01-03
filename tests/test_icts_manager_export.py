from io import BytesIO

import openpyxl
import pytest
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.urls import reverse

from icts.models import AccessProposal


def _create_user(username, groups):
    User = get_user_model()
    user = User.objects.create_user(
        username=username,
        password="safe-pass",
        email=f"{username}@example.com",
    )
    for name in groups:
        group, _ = Group.objects.get_or_create(name=name)
        user.groups.add(group)
    return user


@pytest.mark.django_db
def test_export_manager_summary_excludes_olmat_only_by_default(client):
    manager = _create_user("manager_export", ["managers"])
    applicant = _create_user("export_applicant", ["icts_users"])
    AccessProposal.objects.create(
        applicant=applicant,
        title="OLMAT only export",
        facility_olmat=True,
    )
    AccessProposal.objects.create(
        applicant=applicant,
        title="ICTS export",
        facility_sem=True,
    )

    client.force_login(manager)
    response = client.get(reverse("icts:export_manager_data"))

    assert response.status_code == 200
    workbook = openpyxl.load_workbook(BytesIO(response.content))
    worksheet = workbook.active
    metrics = {
        row[0]: row[1]
        for row in worksheet.iter_rows(min_row=2, values_only=True)
        if row and row[0]
    }
    assert metrics.get("Total propuestas") == 1

    all_text = [
        cell
        for row in worksheet.iter_rows(values_only=True)
        for cell in row
        if isinstance(cell, str)
    ]
    assert "ICTS export" not in all_text
    assert "OLMAT only export" not in all_text


@pytest.mark.django_db
def test_export_manager_rejects_proposals_export_type(client):
    manager = _create_user("manager_export_reject", ["managers"])
    client.force_login(manager)

    response = client.get(f"{reverse('icts:export_manager_data')}?type=proposals")

    assert response.status_code == 400
