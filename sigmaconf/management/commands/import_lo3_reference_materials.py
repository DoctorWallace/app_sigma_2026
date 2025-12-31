from datetime import datetime

from django.core.management.base import BaseCommand, CommandError
from django.utils import timezone

from openpyxl import load_workbook

from sigmaconf.models import LO3ReferenceMaterial


def _parse_date(value):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, str):
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
            try:
                return datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                continue
    return None


def _as_text(value):
    if value is None:
        return ""
    return str(value).strip()


class Command(BaseCommand):
    help = "Importa materiales de referencia LO3 desde Excel (sheet 'LISTADO MR')."

    def add_arguments(self, parser):
        parser.add_argument("--file", required=True, help="Ruta al Excel MR L03-confocal.xlsx")

    def handle(self, *args, **options):
        file_path = options["file"]
        try:
            workbook = load_workbook(filename=file_path, data_only=True)
        except FileNotFoundError as exc:
            raise CommandError(f"No se encuentra el archivo: {file_path}") from exc
        except Exception as exc:
            raise CommandError(f"No se pudo abrir el archivo: {exc}") from exc

        if "LISTADO MR" not in workbook.sheetnames:
            raise CommandError("No se encuentra la hoja 'LISTADO MR'.")

        sheet = workbook["LISTADO MR"]
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            self.stdout.write(self.style.WARNING("La hoja esta vacia."))
            return

        headers = [str(value).strip() if value is not None else "" for value in rows[0]]
        header_map = {header.lower(): idx for idx, header in enumerate(headers)}

        def get_value(row, key):
            idx = header_map.get(key)
            if idx is None:
                return None
            return row[idx]

        imported = 0
        updated = 0
        for row in rows[1:]:
            code = get_value(row, "codigo") or get_value(row, "código")
            if not code:
                continue
            code = _as_text(code)
            defaults = {
                "description": _as_text(
                    get_value(row, "descripcion") or get_value(row, "descripción")
                ),
                "catalog_reference": _as_text(get_value(row, "referencia")),
                "responsible": _as_text(get_value(row, "responsable")),
                "location": _as_text(
                    get_value(row, "localizacion") or get_value(row, "localización")
                ),
                "reception_date": _parse_date(get_value(row, "fecha recepcion") or get_value(row, "fecha recepción")),
                "expiry_date": _parse_date(get_value(row, "caducidad")),
                "observations": _as_text(get_value(row, "observaciones")),
            }

            obj, created = LO3ReferenceMaterial.objects.update_or_create(
                code=code, defaults=defaults
            )
            if created:
                imported += 1
            else:
                updated += 1

            if obj.expiry_date and obj.expiry_date < timezone.localdate():
                if obj.status != "retired":
                    obj.status = "retired"
                    obj.save(update_fields=["status"])

        self.stdout.write(
            self.style.SUCCESS(
                f"Importados: {imported}, actualizados: {updated}."
            )
        )
