import re
import textwrap
from io import BytesIO

from django.contrib import messages
from django.conf import settings
from django.core.files.base import ContentFile
from django.db.models import Case, Count, IntegerField, Max, Prefetch, Q, Value, When
from django.http import FileResponse, Http404, HttpResponse, HttpResponseForbidden
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from datetime import datetime, timedelta

from openpyxl import Workbook

from core.quality_templates import load_quality_json
from icts.decorators import login_required_icts
from icts.auth_utils import (
    CONF_TECH_GROUPS,
    MANAGER_GROUPS,
    RESPONSABLE_GROUPS,
    get_normalized_user_groups,
    user_in_groups,
)
from icts.models import AccessProposal

from .forms import (
    LO3AttachmentForm,
    LO3CalibrationForm,
    LO3IncidentAttachmentForm,
    LO3IncidentForm,
    LO3HabilitationForm,
    LO3MessageForm,
    LO3EquipmentDocumentForm,
    LO3EquipmentForm,
    LO3ReferenceMaterialDocumentForm,
    LO3ReferenceMaterialForm,
    MCFSessionCreateForm,
    MCFFinalReportUploadForm,
    MCFSampleRecordForm,
)
from .models import (
    LO3Attachment,
    LO3Calibration,
    LO3Conversation,
    LO3Message,
    LO3MessageAttachment,
    LO3Equipment,
    LO3EquipmentDocument,
    LO3ReferenceMaterial,
    LO3ReferenceMaterialDocument,
    LO3Incident,
    LO3IncidentAttachment,
    LO3Habilitation,
    MCFSession,
    MCFSampleRecord,
)


def _is_confocal_technician(user):
    groups = get_normalized_user_groups(user)
    return user_in_groups(user, CONF_TECH_GROUPS, groups)


def _ensure_confocal_technician(request):
    if not _is_confocal_technician(request.user):
        messages.error(request, "No tienes permisos para acceder a esta seccion.")
        return False
    return True


def _is_lo3_blocked(user, groups=None):
    if not getattr(user, "is_authenticated", False):
        return False
    if user.is_staff or user.is_superuser:
        return False
    groups = groups or get_normalized_user_groups(user)
    if user_in_groups(user, CONF_TECH_GROUPS, groups):
        return False
    return user_in_groups(user, RESPONSABLE_GROUPS | MANAGER_GROUPS, groups)


def _can_edit_session(user, session):
    groups = get_normalized_user_groups(user)
    if _is_lo3_blocked(user, groups):
        return False
    is_tech = user_in_groups(user, CONF_TECH_GROUPS, groups)
    is_privileged = user.is_staff or user.is_superuser
    return (is_tech and session.technician_id == user.id) or is_privileged


def _can_view_session(user, session):
    groups = get_normalized_user_groups(user)
    if _is_lo3_blocked(user, groups):
        return False, True
    is_tech = user_in_groups(user, CONF_TECH_GROUPS, groups)
    is_applicant = session.access_proposal.applicant_id == user.id
    is_privileged = user.is_staff or user.is_superuser
    allowed = is_applicant or is_tech or is_privileged
    can_edit = (is_tech and session.technician_id == user.id) or is_privileged
    read_only = not can_edit
    return allowed, read_only


def _lo3_access_flags(user):
    groups = get_normalized_user_groups(user)
    is_privileged = user.is_staff or user.is_superuser
    is_tech = user_in_groups(user, CONF_TECH_GROUPS, groups) or is_privileged
    return is_tech, is_privileged


def _lo3_nav_context(user, active_key):
    is_tech, is_privileged = _lo3_access_flags(user)
    return {
        "lo3_is_tech": is_tech,
        "lo3_is_privileged": is_privileged,
        "lo3_active": active_key,
    }


def _is_lo3_principal(user):
    if not getattr(user, "is_authenticated", False):
        return False
    if getattr(user, "is_superuser", False):
        return True
    groups = get_normalized_user_groups(user)
    principal_groups = getattr(
        settings,
        "LO3_PRINCIPAL_GROUPS",
        ["tecnico_responsable_s_conf", "tecnicos_responsables_s_conf"],
    )
    return user_in_groups(user, principal_groups, groups)


def _can_access_lo3_conversation(user, proposal):
    groups = get_normalized_user_groups(user)
    if _is_lo3_blocked(user, groups):
        return False
    is_tech = user_in_groups(user, CONF_TECH_GROUPS, groups)
    is_privileged = user.is_staff or user.is_superuser
    is_applicant = proposal.applicant_id == user.id
    return is_applicant or is_tech or is_privileged


def _calibration_due_status(calibration, today=None):
    if today is None:
        today = timezone.localdate()
    if not calibration.due_at:
        return "no_due", "Sin vencimiento"
    if calibration.due_at < today:
        return "overdue", "Vencida"
    if calibration.due_at <= today + timedelta(days=30):
        return "due_soon", "Vence pronto"
    return "ok", "Al d\u00eda"


def _build_notice_pdf(session):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buffer = BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4, pageCompression=0)
    height = A4[1]

    proposal = session.access_proposal
    access_code = proposal.access_code or f"#{proposal.id}"
    applicant_name = proposal.contact_person or proposal.applicant.get_full_name() or proposal.applicant.get_username()
    profile = getattr(proposal.applicant, "icts_profile", None)
    center = (getattr(profile, "center", "") or "").strip()
    address = (getattr(profile, "address", "") or "").strip()
    completion_date = session.completion_date or timezone.now().date()
    date_str = completion_date.strftime("%d/%m/%Y")

    default_title = "Responsable de LO3 Metrolog\u00eda de superficies \u00f3pticas 3D"
    default_name = "Dra. Teresa Hern\u00e1ndez D\u00edaz"
    responsible_title = getattr(settings, "LO3_RESPONSIBLE_TITLE", default_title) or default_title
    responsible_name = getattr(settings, "LO3_RESPONSIBLE_NAME", default_name) or default_name

    def add_paragraph(text_obj, content):
        for line in textwrap.wrap(content, width=90):
            text_obj.textLine(line)

    text = pdf.beginText(50, height - 70)
    text.setFont("Helvetica", 11)
    text.textLine(f"A la atenci\u00f3n de D/D\u00f1a {applicant_name}")
    if center:
        text.textLine(center)
    if address:
        text.textLine(address)
    text.textLine("")
    text.textLine(f"Madrid, {date_str}")
    text.textLine("")
    text.textLine("Estimado Sr/Sra.:")
    text.textLine("")
    add_paragraph(
        text,
        "Adjunto le enviamos el resultado de las determinaciones efectuadas mediante "
        f"Metrolog\u00eda de superficies \u00f3pticas 3D del lote: {session.lot_code}.",
    )
    text.textLine("")
    add_paragraph(
        text,
        "Puede encontrar dichos informes en su cuenta de usuario SIGMA ICTS, "
        f"asociados a la solicitud {access_code}.",
    )
    text.textLine("")
    text.textLine("Un cordial saludo")
    text.textLine("")
    text.textLine(responsible_title)
    text.textLine(responsible_name)

    pdf.drawText(text)
    pdf.showPage()
    pdf.save()
    buffer.seek(0)
    return buffer.read()


def _parse_confocal_samples(confocal_data):
    if not isinstance(confocal_data, dict):
        return []
    pattern = re.compile(r"^(?:confocal_)?sample_(\d+)_(identification|name|details)$")
    grouped = {}
    for key, value in confocal_data.items():
        match = pattern.match(key)
        if not match:
            continue
        idx = int(match.group(1))
        field = match.group(2)
        grouped.setdefault(idx, {})[field] = value
    return [grouped[idx] for idx in sorted(grouped.keys())]


@login_required_icts
def mcf_dashboard(request):
    groups = get_normalized_user_groups(request.user)
    if _is_lo3_blocked(request.user, groups):
        return HttpResponseForbidden("No autorizado.")
    if not _ensure_confocal_technician(request):
        return redirect("sigmaconf:home")

    accepted_proposals = AccessProposal.objects.filter(
        status="accepted", facility_confocal=True
    ).order_by("-created_at")
    sessions = (
        MCFSession.objects.filter(technician=request.user)
        .annotate(
            status_rank=Case(
                When(status="in_progress", then=Value(0)),
                default=Value(1),
                output_field=IntegerField(),
            ),
            sample_count=Count("samples", distinct=True),
            extra_samples=Count("samples", filter=Q(samples__source="session"), distinct=True),
            analyzed_samples=Count(
                "samples", filter=Q(samples__analysis_date__isnull=False), distinct=True
            ),
        )
        .order_by("status_rank", "-created_at")
    )
    context = {
        "accepted_proposals": accepted_proposals,
        "sessions": sessions,
        "in_progress_count": sessions.filter(status="in_progress").count(),
        "completed_count": sessions.filter(status="completed").count(),
    }
    context.update(_lo3_nav_context(request.user, "requests"))
    return render(request, "sigmaconf/mcf_dashboard.html", context)


@login_required_icts
def mcf_create_session(request, proposal_id=None):
    groups = get_normalized_user_groups(request.user)
    if _is_lo3_blocked(request.user, groups):
        return HttpResponseForbidden("No autorizado.")
    if not _ensure_confocal_technician(request):
        return redirect("sigmaconf:home")

    proposal = None
    if proposal_id:
        proposal = get_object_or_404(
            AccessProposal, pk=proposal_id, status="accepted", facility_confocal=True
        )

    if request.method == "POST":
        if proposal is None:
            form = MCFSessionCreateForm(request.POST)
            if not form.is_valid():
                return render(
                    request,
                    "sigmaconf/mcf_create_session.html",
                    {"form": form, "available_proposals": form.fields["access_proposal"].queryset},
                )
            proposal = form.cleaned_data["access_proposal"]

        session = MCFSession(access_proposal=proposal, technician=request.user)
        proposal_data = proposal.facility_data or {}
        if isinstance(proposal_data, dict):
            session.request_snapshot = proposal_data.get("confocal", {}) or {}
        session.save()

        samples_data = _parse_confocal_samples(session.request_snapshot)
        created = 0
        sequence = 1
        for sample in samples_data:
            identification = str(sample.get("identification") or "").strip()
            name = str(sample.get("name") or "").strip()
            details = str(sample.get("details") or "").strip()
            if not identification and not name:
                continue
            MCFSampleRecord.objects.create(
                session=session,
                sequence=sequence,
                source="proposal",
                identification=identification or name or f"Sample {sequence}",
                name=name,
                details=details,
            )
            sequence += 1
            created += 1

        if created == 0:
            MCFSampleRecord.objects.create(
                session=session,
                sequence=1,
                source="proposal",
                identification="Sample 1",
                name="",
                details="",
            )
            messages.warning(request, "Sesion creada con una muestra base sin datos importados.")
        else:
            messages.success(request, f"Sesion creada con {created} muestras importadas.")
        return redirect("sigmaconf:mcf_session_detail", pk=session.pk)
    else:
        form = MCFSessionCreateForm(initial={"access_proposal": proposal} if proposal else None)
        if proposal:
            form.fields["access_proposal"].queryset = AccessProposal.objects.filter(pk=proposal.pk)

    available = AccessProposal.objects.filter(
        status="accepted", facility_confocal=True
    ).order_by("-created_at")
    return render(
        request,
        "sigmaconf/mcf_create_session.html",
        {"form": form, "available_proposals": available},
    )


@login_required_icts
def lo3_home(request):
    groups = get_normalized_user_groups(request.user)
    if _is_lo3_blocked(request.user, groups):
        return HttpResponseForbidden("No autorizado.")
    is_tech, _ = _lo3_access_flags(request.user)
    if not is_tech:
        return redirect("sigmaconf:lo3_results")

    sessions = MCFSession.objects.filter(technician=request.user)
    today = timezone.localdate()
    window_30 = today - timedelta(days=30)
    window_90 = today - timedelta(days=90)

    in_progress_30 = sessions.filter(
        status="in_progress", created_at__date__gte=window_30
    ).count()
    completed_30 = sessions.filter(
        status="completed", completion_date__gte=window_30
    ).count()
    in_progress_90 = sessions.filter(
        status="in_progress", created_at__date__gte=window_90
    ).count()
    completed_90 = sessions.filter(
        status="completed", completion_date__gte=window_90
    ).count()

    completed_sessions = sessions.filter(status="completed")
    completed_total = completed_sessions.count()
    completed_with_report = (
        completed_sessions.exclude(final_report_pdf="")
        .exclude(final_report_pdf__isnull=True)
        .count()
    )
    final_report_rate = (
        round((completed_with_report / completed_total) * 100, 1)
        if completed_total
        else None
    )

    analysis_delays = []
    for session in (
        sessions.select_related("access_proposal").prefetch_related("samples")
    ):
        proposal = session.access_proposal
        base_dt = (
            getattr(proposal, "accepted_at", None)
            or proposal.submitted_at
            or proposal.created_at
        )
        if not base_dt:
            continue
        if isinstance(base_dt, datetime):
            base_date = timezone.localdate(base_dt)
        else:
            base_date = base_dt
        analysis_dates = [
            sample.analysis_date
            for sample in session.samples.all()
            if sample.analysis_date
        ]
        if not analysis_dates:
            continue
        first_analysis = min(analysis_dates)
        delta_days = (first_analysis - base_date).days
        if delta_days >= 0:
            analysis_delays.append(delta_days)

    avg_analysis_delay = (
        round(sum(analysis_delays) / len(analysis_delays), 1)
        if analysis_delays
        else None
    )

    context = {
        "in_progress_count": sessions.filter(status="in_progress").count(),
        "completed_count": sessions.filter(status="completed").count(),
        "accepted_count": AccessProposal.objects.filter(
            status="accepted", facility_confocal=True
        ).count(),
        "total_sessions": sessions.count(),
        "stats": {
            "in_progress_30": in_progress_30,
            "completed_30": completed_30,
            "in_progress_90": in_progress_90,
            "completed_90": completed_90,
            "avg_analysis_delay": avg_analysis_delay,
            "final_report_rate": final_report_rate,
            "completed_total": completed_total,
            "completed_with_report": completed_with_report,
        },
    }
    context.update(_lo3_nav_context(request.user, "home"))
    return render(request, "sigmaconf/lo3_home.html", context)


@login_required_icts
def lo3_results(request):
    groups = get_normalized_user_groups(request.user)
    if _is_lo3_blocked(request.user, groups):
        return HttpResponseForbidden("No autorizado.")
    is_tech, is_privileged = _lo3_access_flags(request.user)
    sessions = MCFSession.objects.filter(status="completed")
    if is_privileged:
        pass
    elif is_tech:
        sessions = sessions.filter(technician=request.user)
    else:
        sessions = sessions.filter(access_proposal__applicant=request.user)

    sessions = (
        sessions.annotate(
            sample_count=Count("samples", distinct=True),
            extra_samples=Count("samples", filter=Q(samples__source="session"), distinct=True),
        )
        .order_by("-completion_date", "-created_at")
    )
    context = {"sessions": sessions}
    context.update(_lo3_nav_context(request.user, "results"))
    return render(request, "sigmaconf/lo3_results.html", context)


@login_required_icts
def lo3_communications_list(request):
    groups = get_normalized_user_groups(request.user)
    if _is_lo3_blocked(request.user, groups):
        return HttpResponseForbidden("No autorizado.")
    is_tech, is_privileged = _lo3_access_flags(request.user)
    is_staff_side = is_tech or is_privileged

    proposals = AccessProposal.objects.filter(
        status="accepted", facility_confocal=True
    ).select_related("applicant")
    if not is_staff_side:
        proposals = proposals.filter(applicant=request.user)
    proposals = proposals.order_by("-created_at")

    conversations = (
        LO3Conversation.objects.filter(access_proposal__in=proposals)
        .select_related("access_proposal", "access_proposal__applicant")
        .prefetch_related(
            Prefetch(
                "messages",
                queryset=LO3Message.objects.select_related("sender").order_by("-created_at"),
            )
        )
    )
    conversation_map = {conversation.access_proposal_id: conversation for conversation in conversations}

    entries = []
    for proposal in proposals:
        conversation = conversation_map.get(proposal.id)
        last_message = None
        unread_count = 0
        if conversation:
            message_list = list(conversation.messages.all())
            if message_list:
                last_message = message_list[0]
                if is_staff_side:
                    unread_count = sum(
                        1 for msg in message_list if msg.read_at_technician is None
                    )
                else:
                    unread_count = sum(1 for msg in message_list if msg.read_at_applicant is None)
        entries.append(
            {
                "proposal": proposal,
                "conversation": conversation,
                "last_message": last_message,
                "unread_count": unread_count,
            }
        )

    context = {"entries": entries, "is_staff_side": is_staff_side}
    context.update(_lo3_nav_context(request.user, "communications"))
    return render(request, "sigmaconf/lo3_communications_list.html", context)


@login_required_icts
def lo3_communications_detail(request, proposal_id):
    groups = get_normalized_user_groups(request.user)
    if _is_lo3_blocked(request.user, groups):
        return HttpResponseForbidden("No autorizado.")
    proposal = get_object_or_404(
        AccessProposal, pk=proposal_id, status="accepted", facility_confocal=True
    )
    if not _can_access_lo3_conversation(request.user, proposal):
        return HttpResponseForbidden("No autorizado.")

    is_tech, is_privileged = _lo3_access_flags(request.user)
    is_staff_side = is_tech or is_privileged
    is_applicant = proposal.applicant_id == request.user.id

    conversation, _ = LO3Conversation.objects.get_or_create(access_proposal=proposal)

    if request.method == "POST":
        form = LO3MessageForm(request.POST, request.FILES)
        if form.is_valid():
            body = (form.cleaned_data.get("body") or "").strip()
            if not body:
                messages.error(request, "El mensaje no puede estar vacío.")
            else:
                message = LO3Message.objects.create(
                    conversation=conversation,
                    sender=request.user,
                    body=body,
                    read_at_applicant=timezone.now() if is_applicant else None,
                    read_at_technician=timezone.now() if is_staff_side else None,
                )
                for file_obj in request.FILES.getlist("attachments"):
                    LO3MessageAttachment.objects.create(
                        message=message,
                        file=file_obj,
                        filename=file_obj.name,
                        uploaded_by=request.user,
                    )
                conversation.save(update_fields=["updated_at"])
                messages.success(request, "Mensaje enviado.")
                return redirect("sigmaconf:lo3_communications_detail", proposal_id=proposal.pk)
    else:
        form = LO3MessageForm()

    messages_qs = (
        conversation.messages.select_related("sender")
        .prefetch_related("attachments")
        .order_by("created_at")
    )

    if is_staff_side:
        unread_qs = messages_qs.filter(read_at_technician__isnull=True)
        if unread_qs.exists():
            unread_qs.update(read_at_technician=timezone.now())
    else:
        unread_qs = messages_qs.filter(read_at_applicant__isnull=True)
        if unread_qs.exists():
            unread_qs.update(read_at_applicant=timezone.now())

    context = {
        "proposal": proposal,
        "conversation": conversation,
        "messages_list": messages_qs,
        "form": form,
        "is_staff_side": is_staff_side,
    }
    context.update(_lo3_nav_context(request.user, "communications"))
    return render(request, "sigmaconf/lo3_communications_detail.html", context)


@login_required_icts
def lo3_message_attachment_download(request, attachment_id):
    attachment = get_object_or_404(LO3MessageAttachment, pk=attachment_id)
    proposal = attachment.message.conversation.access_proposal
    if not _can_access_lo3_conversation(request.user, proposal):
        return HttpResponseForbidden("No autorizado.")
    filename = attachment.filename or attachment.file.name.rsplit("/", 1)[-1]
    return FileResponse(attachment.file.open("rb"), as_attachment=True, filename=filename)


@login_required_icts
def lo3_equipment_list(request):
    is_tech, _ = _lo3_access_flags(request.user)
    if not is_tech:
        return HttpResponseForbidden("No autorizado.")

    equipment_list = LO3Equipment.objects.all()
    context = {"equipment_list": equipment_list}
    context.update(_lo3_nav_context(request.user, "equipment"))
    return render(request, "sigmaconf/lo3_equipment_list.html", context)


@login_required_icts
def lo3_equipment_detail(request, equipment_id):
    is_tech, _ = _lo3_access_flags(request.user)
    if not is_tech:
        return HttpResponseForbidden("No autorizado.")

    equipment = get_object_or_404(LO3Equipment, pk=equipment_id)
    documents = equipment.documents.select_related("uploaded_by")
    context = {
        "equipment": equipment,
        "documents": documents,
        "can_edit": is_tech,
        "document_form": LO3EquipmentDocumentForm(),
    }
    context.update(_lo3_nav_context(request.user, "equipment"))
    return render(request, "sigmaconf/lo3_equipment_detail.html", context)


@login_required_icts
def lo3_equipment_create(request):
    is_tech, _ = _lo3_access_flags(request.user)
    if not is_tech:
        return HttpResponseForbidden("No autorizado.")

    if request.method == "POST":
        form = LO3EquipmentForm(request.POST)
        if form.is_valid():
            equipment = form.save()
            messages.success(request, "Equipo creado.")
            return redirect("sigmaconf:lo3_equipment_detail", equipment_id=equipment.pk)
    else:
        form = LO3EquipmentForm()

    context = {"form": form, "is_editing": False}
    context.update(_lo3_nav_context(request.user, "equipment"))
    return render(request, "sigmaconf/lo3_equipment_form.html", context)


@login_required_icts
def lo3_equipment_edit(request, equipment_id):
    is_tech, _ = _lo3_access_flags(request.user)
    if not is_tech:
        return HttpResponseForbidden("No autorizado.")

    equipment = get_object_or_404(LO3Equipment, pk=equipment_id)
    if request.method == "POST":
        form = LO3EquipmentForm(request.POST, instance=equipment)
        if form.is_valid():
            form.save()
            messages.success(request, "Equipo actualizado.")
            return redirect("sigmaconf:lo3_equipment_detail", equipment_id=equipment.pk)
    else:
        form = LO3EquipmentForm(instance=equipment)

    context = {"form": form, "equipment": equipment, "is_editing": True}
    context.update(_lo3_nav_context(request.user, "equipment"))
    return render(request, "sigmaconf/lo3_equipment_form.html", context)


@login_required_icts
def lo3_equipment_delete(request, equipment_id):
    if request.method != "POST":
        return HttpResponseForbidden("Metodo no permitido.")

    is_tech, _ = _lo3_access_flags(request.user)
    if not is_tech:
        return HttpResponseForbidden("No autorizado.")

    equipment = get_object_or_404(LO3Equipment, pk=equipment_id)
    equipment.delete()
    messages.success(request, "Equipo eliminado.")
    return redirect("sigmaconf:lo3_equipment_list")


@login_required_icts
def lo3_equipment_document_add(request, equipment_id):
    if request.method != "POST":
        return HttpResponseForbidden("Metodo no permitido.")

    is_tech, _ = _lo3_access_flags(request.user)
    if not is_tech:
        return HttpResponseForbidden("No autorizado.")

    equipment = get_object_or_404(LO3Equipment, pk=equipment_id)
    form = LO3EquipmentDocumentForm(request.POST, request.FILES)
    if form.is_valid():
        document = form.save(commit=False)
        document.equipment = equipment
        document.uploaded_by = request.user
        document.save()
        messages.success(request, "Documento subido.")
    else:
        messages.error(request, "No se pudo subir el documento.")
    return redirect("sigmaconf:lo3_equipment_detail", equipment_id=equipment.pk)


@login_required_icts
def lo3_equipment_document_download(request, document_id):
    is_tech, _ = _lo3_access_flags(request.user)
    if not is_tech:
        return HttpResponseForbidden("No autorizado.")

    document = get_object_or_404(LO3EquipmentDocument, pk=document_id)
    filename = document.file.name.rsplit("/", 1)[-1]
    return FileResponse(document.file.open("rb"), as_attachment=True, filename=filename)


@login_required_icts
def lo3_equipment_document_delete(request, document_id):
    if request.method != "POST":
        return HttpResponseForbidden("Metodo no permitido.")

    is_tech, _ = _lo3_access_flags(request.user)
    if not is_tech:
        return HttpResponseForbidden("No autorizado.")

    document = get_object_or_404(LO3EquipmentDocument, pk=document_id)
    equipment_id = document.equipment_id
    document.delete()
    messages.success(request, "Documento eliminado.")
    return redirect("sigmaconf:lo3_equipment_detail", equipment_id=equipment_id)


def _get_calibration_type(calibration_type):
    mapping = {
        "interna": ("internal", "Calibraciones internas"),
        "externa": ("external", "Calibraciones externas"),
    }
    return mapping.get(calibration_type)


@login_required_icts
def lo3_calibration_list(request, calibration_type):
    is_tech, _ = _lo3_access_flags(request.user)
    if not is_tech:
        return HttpResponseForbidden("No autorizado.")

    type_info = _get_calibration_type(calibration_type)
    if not type_info:
        raise Http404("Tipo de calibracion no valido.")
    type_key, title = type_info

    calibrations = (
        LO3Calibration.objects.filter(calibration_type=type_key)
        .select_related("equipment")
        .order_by("-performed_at")
    )
    today = timezone.localdate()
    for calibration in calibrations:
        status, label = _calibration_due_status(calibration, today=today)
        calibration.due_status = status
        calibration.due_label = label

    context = {
        "calibrations": calibrations,
        "calibration_type": calibration_type,
        "title": title,
    }
    context.update(_lo3_nav_context(request.user, "calibrations"))
    return render(request, "sigmaconf/lo3_calibration_list.html", context)


@login_required_icts
def lo3_calibration_create(request, calibration_type):
    is_tech, _ = _lo3_access_flags(request.user)
    if not is_tech:
        return HttpResponseForbidden("No autorizado.")

    type_info = _get_calibration_type(calibration_type)
    if not type_info:
        raise Http404("Tipo de calibracion no valido.")
    type_key, title = type_info

    if request.method == "POST":
        form = LO3CalibrationForm(request.POST, request.FILES)
        if form.is_valid():
            calibration = form.save(commit=False)
            calibration.calibration_type = type_key
            calibration.save()
            messages.success(request, "Calibracion registrada.")
            return redirect(
                "sigmaconf:lo3_calibration_list", calibration_type=calibration_type
            )
    else:
        form = LO3CalibrationForm()

    context = {
        "form": form,
        "is_editing": False,
        "calibration_type": calibration_type,
        "title": title,
    }
    context.update(_lo3_nav_context(request.user, "calibrations"))
    return render(request, "sigmaconf/lo3_calibration_form.html", context)


@login_required_icts
def lo3_calibration_edit(request, calibration_id):
    is_tech, _ = _lo3_access_flags(request.user)
    if not is_tech:
        return HttpResponseForbidden("No autorizado.")

    calibration = get_object_or_404(LO3Calibration, pk=calibration_id)
    calibration_type = "interna" if calibration.calibration_type == "internal" else "externa"
    type_info = _get_calibration_type(calibration_type)
    title = type_info[1] if type_info else "Calibraciones"

    if request.method == "POST":
        form = LO3CalibrationForm(request.POST, request.FILES, instance=calibration)
        if form.is_valid():
            form.save()
            messages.success(request, "Calibracion actualizada.")
            return redirect(
                "sigmaconf:lo3_calibration_list", calibration_type=calibration_type
            )
    else:
        form = LO3CalibrationForm(instance=calibration)

    context = {
        "form": form,
        "is_editing": True,
        "calibration": calibration,
        "calibration_type": calibration_type,
        "title": title,
    }
    context.update(_lo3_nav_context(request.user, "calibrations"))
    return render(request, "sigmaconf/lo3_calibration_form.html", context)


@login_required_icts
def lo3_calibration_delete(request, calibration_id):
    if request.method != "POST":
        return HttpResponseForbidden("Metodo no permitido.")

    is_tech, _ = _lo3_access_flags(request.user)
    if not is_tech:
        return HttpResponseForbidden("No autorizado.")

    calibration = get_object_or_404(LO3Calibration, pk=calibration_id)
    calibration_type = "interna" if calibration.calibration_type == "internal" else "externa"
    calibration.delete()
    messages.success(request, "Calibracion eliminada.")
    return redirect("sigmaconf:lo3_calibration_list", calibration_type=calibration_type)


@login_required_icts
def lo3_calibration_certificate_download(request, calibration_id):
    is_tech, _ = _lo3_access_flags(request.user)
    if not is_tech:
        return HttpResponseForbidden("No autorizado.")

    calibration = get_object_or_404(LO3Calibration, pk=calibration_id)
    if not calibration.certificate_file:
        raise Http404("Certificado no disponible.")
    filename = calibration.certificate_file.name.rsplit("/", 1)[-1]
    return FileResponse(
        calibration.certificate_file.open("rb"), as_attachment=True, filename=filename
    )


def _reference_material_status(material, today=None):
    if today is None:
        today = timezone.localdate()
    if material.status == "retired":
        return "retired", "Retirado"
    if material.expiry_date and material.expiry_date < today:
        return "expired", "Caducado"
    if material.expiry_date and material.expiry_date <= today + timedelta(days=30):
        return "due_soon", "Caduca pronto"
    return "active", "Activo"


@login_required_icts
def lo3_reference_material_list(request):
    is_tech, _ = _lo3_access_flags(request.user)
    if not is_tech:
        return HttpResponseForbidden("No autorizado.")

    materials = LO3ReferenceMaterial.objects.all().order_by("code")
    today = timezone.localdate()
    expired_count = 0
    active_count = 0
    for material in materials:
        status, label = _reference_material_status(material, today=today)
        material.status_label = label
        material.status_style = status
        if status in ("expired", "due_soon"):
            expired_count += 1
        elif status == "active":
            active_count += 1

    context = {
        "materials": materials,
        "expired_count": expired_count,
        "active_count": active_count,
    }
    context.update(_lo3_nav_context(request.user, "materials"))
    return render(request, "sigmaconf/lo3_reference_material_list.html", context)


@login_required_icts
def lo3_reference_material_detail(request, material_id):
    is_tech, _ = _lo3_access_flags(request.user)
    if not is_tech:
        return HttpResponseForbidden("No autorizado.")

    material = get_object_or_404(LO3ReferenceMaterial, pk=material_id)
    documents = material.documents.all()
    context = {
        "material": material,
        "documents": documents,
        "document_form": LO3ReferenceMaterialDocumentForm(),
    }
    context.update(_lo3_nav_context(request.user, "materials"))
    return render(request, "sigmaconf/lo3_reference_material_detail.html", context)


@login_required_icts
def lo3_reference_material_create(request):
    is_tech, _ = _lo3_access_flags(request.user)
    if not is_tech:
        return HttpResponseForbidden("No autorizado.")

    if request.method == "POST":
        form = LO3ReferenceMaterialForm(request.POST)
        if form.is_valid():
            material = form.save()
            messages.success(request, "Material creado.")
            return redirect(
                "sigmaconf:lo3_reference_material_detail", material_id=material.pk
            )
    else:
        form = LO3ReferenceMaterialForm()

    context = {"form": form, "is_editing": False}
    context.update(_lo3_nav_context(request.user, "materials"))
    return render(request, "sigmaconf/lo3_reference_material_form.html", context)


@login_required_icts
def lo3_reference_material_edit(request, material_id):
    is_tech, _ = _lo3_access_flags(request.user)
    if not is_tech:
        return HttpResponseForbidden("No autorizado.")

    material = get_object_or_404(LO3ReferenceMaterial, pk=material_id)
    if request.method == "POST":
        form = LO3ReferenceMaterialForm(request.POST, instance=material)
        if form.is_valid():
            form.save()
            messages.success(request, "Material actualizado.")
            return redirect(
                "sigmaconf:lo3_reference_material_detail", material_id=material.pk
            )
    else:
        form = LO3ReferenceMaterialForm(instance=material)

    context = {"form": form, "is_editing": True, "material": material}
    context.update(_lo3_nav_context(request.user, "materials"))
    return render(request, "sigmaconf/lo3_reference_material_form.html", context)


@login_required_icts
def lo3_reference_material_delete(request, material_id):
    if request.method != "POST":
        return HttpResponseForbidden("Metodo no permitido.")

    is_tech, _ = _lo3_access_flags(request.user)
    if not is_tech:
        return HttpResponseForbidden("No autorizado.")

    material = get_object_or_404(LO3ReferenceMaterial, pk=material_id)
    material.delete()
    messages.success(request, "Material eliminado.")
    return redirect("sigmaconf:lo3_reference_material_list")


@login_required_icts
def lo3_reference_material_document_add(request, material_id):
    if request.method != "POST":
        return HttpResponseForbidden("Metodo no permitido.")

    is_tech, _ = _lo3_access_flags(request.user)
    if not is_tech:
        return HttpResponseForbidden("No autorizado.")

    material = get_object_or_404(LO3ReferenceMaterial, pk=material_id)
    form = LO3ReferenceMaterialDocumentForm(request.POST, request.FILES)
    if form.is_valid():
        document = form.save(commit=False)
        document.material = material
        document.save()
        messages.success(request, "Documento subido.")
    else:
        messages.error(request, "No se pudo subir el documento.")
    return redirect("sigmaconf:lo3_reference_material_detail", material_id=material.pk)


@login_required_icts
def lo3_reference_material_document_download(request, document_id):
    is_tech, _ = _lo3_access_flags(request.user)
    if not is_tech:
        return HttpResponseForbidden("No autorizado.")

    document = get_object_or_404(LO3ReferenceMaterialDocument, pk=document_id)
    filename = document.file.name.rsplit("/", 1)[-1]
    return FileResponse(document.file.open("rb"), as_attachment=True, filename=filename)


@login_required_icts
def lo3_reference_material_document_delete(request, document_id):
    if request.method != "POST":
        return HttpResponseForbidden("Metodo no permitido.")

    is_tech, _ = _lo3_access_flags(request.user)
    if not is_tech:
        return HttpResponseForbidden("No autorizado.")

    document = get_object_or_404(LO3ReferenceMaterialDocument, pk=document_id)
    material_id = document.material_id
    document.delete()
    messages.success(request, "Documento eliminado.")
    return redirect("sigmaconf:lo3_reference_material_detail", material_id=material_id)


def _coerce_incident_choice(value, choices):
    allowed = {key for key, _label in choices}
    return value if value in allowed else None


@login_required_icts
def lo3_incident_list(request):
    is_tech, _ = _lo3_access_flags(request.user)
    if not is_tech:
        return HttpResponseForbidden("No autorizado.")

    incidents = LO3Incident.objects.select_related("equipment", "session").order_by(
        "-created_at"
    )
    incident_type = _coerce_incident_choice(
        request.GET.get("type"), LO3Incident.INCIDENT_TYPE_CHOICES
    )
    status = _coerce_incident_choice(request.GET.get("status"), LO3Incident.STATUS_CHOICES)
    severity = _coerce_incident_choice(
        request.GET.get("severity"), LO3Incident.SEVERITY_CHOICES
    )
    if incident_type:
        incidents = incidents.filter(incident_type=incident_type)
    if status:
        incidents = incidents.filter(status=status)
    if severity:
        incidents = incidents.filter(severity=severity)

    context = {
        "incidents": incidents,
        "filters": {
            "type": incident_type or "",
            "status": status or "",
            "severity": severity or "",
        },
    }
    context.update(_lo3_nav_context(request.user, "incidences"))
    return render(request, "sigmaconf/lo3_incident_list.html", context)


@login_required_icts
def lo3_incident_detail(request, incident_id):
    is_tech, _ = _lo3_access_flags(request.user)
    if not is_tech:
        return HttpResponseForbidden("No autorizado.")

    incident = get_object_or_404(LO3Incident, pk=incident_id)
    attachments = incident.attachments.select_related("uploaded_by")
    context = {
        "incident": incident,
        "attachments": attachments,
        "attachment_form": LO3IncidentAttachmentForm(),
    }
    context.update(_lo3_nav_context(request.user, "incidences"))
    return render(request, "sigmaconf/lo3_incident_detail.html", context)


def _sync_incident_closed_at(incident):
    if incident.status == "closed":
        if not incident.closed_at:
            incident.closed_at = timezone.now()
    else:
        incident.closed_at = None


@login_required_icts
def lo3_incident_create(request):
    is_tech, _ = _lo3_access_flags(request.user)
    if not is_tech:
        return HttpResponseForbidden("No autorizado.")

    if request.method == "POST":
        form = LO3IncidentForm(request.POST)
        if form.is_valid():
            incident = form.save(commit=False)
            incident.created_by = request.user
            _sync_incident_closed_at(incident)
            incident.save()
            messages.success(request, "Incidencia registrada.")
            return redirect("sigmaconf:lo3_incident_detail", incident_id=incident.pk)
    else:
        form = LO3IncidentForm()

    context = {"form": form, "is_editing": False}
    context.update(_lo3_nav_context(request.user, "incidences"))
    return render(request, "sigmaconf/lo3_incident_form.html", context)


@login_required_icts
def lo3_incident_edit(request, incident_id):
    is_tech, _ = _lo3_access_flags(request.user)
    if not is_tech:
        return HttpResponseForbidden("No autorizado.")

    incident = get_object_or_404(LO3Incident, pk=incident_id)
    if request.method == "POST":
        form = LO3IncidentForm(request.POST, instance=incident)
        if form.is_valid():
            incident = form.save(commit=False)
            _sync_incident_closed_at(incident)
            incident.save()
            messages.success(request, "Incidencia actualizada.")
            return redirect("sigmaconf:lo3_incident_detail", incident_id=incident.pk)
    else:
        form = LO3IncidentForm(instance=incident)

    context = {"form": form, "incident": incident, "is_editing": True}
    context.update(_lo3_nav_context(request.user, "incidences"))
    return render(request, "sigmaconf/lo3_incident_form.html", context)


@login_required_icts
def lo3_incident_close(request, incident_id):
    if request.method != "POST":
        return HttpResponseForbidden("Metodo no permitido.")

    is_tech, _ = _lo3_access_flags(request.user)
    if not is_tech:
        return HttpResponseForbidden("No autorizado.")

    incident = get_object_or_404(LO3Incident, pk=incident_id)
    incident.status = "closed"
    _sync_incident_closed_at(incident)
    incident.save(update_fields=["status", "closed_at"])
    messages.success(request, "Incidencia cerrada.")
    return redirect("sigmaconf:lo3_incident_detail", incident_id=incident.pk)


@login_required_icts
def lo3_incident_attachment_add(request, incident_id):
    if request.method != "POST":
        return HttpResponseForbidden("Metodo no permitido.")

    is_tech, _ = _lo3_access_flags(request.user)
    if not is_tech:
        return HttpResponseForbidden("No autorizado.")

    incident = get_object_or_404(LO3Incident, pk=incident_id)
    form = LO3IncidentAttachmentForm(request.POST, request.FILES)
    if form.is_valid():
        attachment = form.save(commit=False)
        attachment.incident = incident
        attachment.uploaded_by = request.user
        attachment.save()
        messages.success(request, "Adjunto subido.")
    else:
        messages.error(request, "No se pudo subir el adjunto.")
    return redirect("sigmaconf:lo3_incident_detail", incident_id=incident.pk)


@login_required_icts
def lo3_incident_attachment_download(request, attachment_id):
    is_tech, _ = _lo3_access_flags(request.user)
    if not is_tech:
        return HttpResponseForbidden("No autorizado.")

    attachment = get_object_or_404(LO3IncidentAttachment, pk=attachment_id)
    filename = attachment.file.name.rsplit("/", 1)[-1]
    return FileResponse(attachment.file.open("rb"), as_attachment=True, filename=filename)


@login_required_icts
def lo3_incident_attachment_delete(request, attachment_id):
    if request.method != "POST":
        return HttpResponseForbidden("Metodo no permitido.")

    is_tech, _ = _lo3_access_flags(request.user)
    if not is_tech:
        return HttpResponseForbidden("No autorizado.")

    attachment = get_object_or_404(LO3IncidentAttachment, pk=attachment_id)
    incident_id = attachment.incident_id
    attachment.delete()
    messages.success(request, "Adjunto eliminado.")
    return redirect("sigmaconf:lo3_incident_detail", incident_id=incident_id)


def _habilitated_name(habilitation):
    if habilitation.investigator_full_name:
        return habilitation.investigator_full_name
    if habilitation.investigator_user:
        return (
            habilitation.investigator_user.get_full_name()
            or habilitation.investigator_user.username
        )
    return ""


def _build_habilitation_pdf(habilitation):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buffer = BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4, pageCompression=0)
    height = A4[1]

    default_title = "Responsable de LO3 Metrolog\u00eda de superficies \u00f3pticas 3D"
    default_name = "Dra. Teresa Hern\u00e1ndez D\u00edaz"
    responsible_title = getattr(settings, "LO3_RESPONSIBLE_TITLE", default_title) or default_title
    responsible_name = getattr(settings, "LO3_RESPONSIBLE_NAME", default_name) or default_name

    issued_at = habilitation.issued_at or timezone.localdate()
    issued_str = issued_at.strftime("%d/%m/%Y")
    valid_until = habilitation.valid_until
    valid_str = valid_until.strftime("%d/%m/%Y") if valid_until else "Sin fecha de caducidad"

    person_name = _habilitated_name(habilitation)
    ciemat_id = habilitation.ciemat_id or "N/D"
    department = habilitation.department_division or "N/D"
    center = habilitation.center_university or "N/D"
    equipment = habilitation.equipment
    equipment_label = equipment.name
    if equipment.manufacturer or equipment.model:
        equipment_label = f"{equipment.name} ({equipment.manufacturer} {equipment.model})".strip()

    def add_paragraph(text_obj, content):
        for line in textwrap.wrap(content, width=92):
            text_obj.textLine(line)

    text = pdf.beginText(50, height - 70)
    text.setFont("Helvetica-Bold", 14)
    text.textLine("CERTIFICADO DE HABILITACION")
    text.setFont("Helvetica", 11)
    text.textLine("")
    text.textLine(f"Madrid, {issued_str}")
    text.textLine("")
    add_paragraph(
        text,
        f"El/la abajo firmante, {responsible_name}, {responsible_title}, "
        "certifica que:",
    )
    text.textLine("")
    add_paragraph(
        text,
        f"D./D\u00f1a {person_name}, con identificador {ciemat_id}, "
        f"perteneciente a {department} ({center}), ha recibido la formaci\u00f3n "
        "necesaria para el uso aut\u00f3nomo del equipo siguiente:",
    )
    text.textLine("")
    add_paragraph(text, f"- {equipment_label}.")
    text.textLine("")
    add_paragraph(
        text,
        "Compromisos: cumplir los procedimientos de seguridad, respetar los "
        "protocolos de uso y registrar cualquier incidencia durante el manejo.",
    )
    text.textLine("")
    add_paragraph(text, f"Validez del certificado: {valid_str}.")
    text.textLine("")
    text.textLine("Y para que conste, se firma el presente certificado.")
    text.textLine("")
    text.textLine(responsible_title)
    text.textLine(responsible_name)

    pdf.drawText(text)
    pdf.showPage()
    pdf.save()
    buffer.seek(0)
    return buffer.read()


def _can_view_habilitation(user, habilitation):
    groups = get_normalized_user_groups(user)
    if _is_lo3_blocked(user, groups):
        return False
    is_tech = user_in_groups(user, CONF_TECH_GROUPS, groups)
    is_privileged = user.is_staff or user.is_superuser
    is_investigator = habilitation.investigator_user_id == user.id
    is_linked = habilitation.linked_technician_account_id == user.id
    return is_tech or is_privileged or is_investigator or is_linked


@login_required_icts
def lo3_habilitation_list(request):
    is_tech, is_privileged = _lo3_access_flags(request.user)
    if not (is_tech or is_privileged):
        return HttpResponseForbidden("No autorizado.")

    habilitations = (
        LO3Habilitation.objects.select_related("equipment", "investigator_user")
        .order_by("-issued_at")
    )
    context = {"habilitations": habilitations, "is_principal": _is_lo3_principal(request.user)}
    context.update(_lo3_nav_context(request.user, "habilitations"))
    return render(request, "sigmaconf/lo3_habilitation_list.html", context)


@login_required_icts
def lo3_habilitation_detail(request, habilitation_id):
    habilitation = get_object_or_404(LO3Habilitation, pk=habilitation_id)
    if not _can_view_habilitation(request.user, habilitation):
        return HttpResponseForbidden("No autorizado.")

    context = {
        "habilitation": habilitation,
        "can_revoke": _is_lo3_principal(request.user),
        "person_name": _habilitated_name(habilitation),
    }
    context.update(_lo3_nav_context(request.user, "habilitations"))
    return render(request, "sigmaconf/lo3_habilitation_detail.html", context)


@login_required_icts
def lo3_habilitation_create(request):
    if not _is_lo3_principal(request.user):
        return HttpResponseForbidden("No autorizado.")

    if request.method == "POST":
        form = LO3HabilitationForm(request.POST)
        if form.is_valid():
            habilitation = form.save(commit=False)
            habilitation.issued_by = request.user
            if not habilitation.issued_at:
                habilitation.issued_at = timezone.localdate()
            pdf_bytes = _build_habilitation_pdf(habilitation)
            filename = f"Habilitacion_LO3_{timezone.now().strftime('%Y%m%d%H%M%S')}.pdf"
            habilitation.certificate_pdf.save(filename, ContentFile(pdf_bytes), save=False)
            habilitation.save()
            messages.success(request, "Habilitacion generada correctamente.")
            return redirect(
                "sigmaconf:lo3_habilitation_detail", habilitation_id=habilitation.pk
            )
    else:
        form = LO3HabilitationForm(initial={"issued_at": timezone.localdate()})

    context = {"form": form}
    context.update(_lo3_nav_context(request.user, "habilitations"))
    return render(request, "sigmaconf/lo3_habilitation_form.html", context)


@login_required_icts
def lo3_habilitation_download(request, habilitation_id):
    habilitation = get_object_or_404(LO3Habilitation, pk=habilitation_id)
    if not _can_view_habilitation(request.user, habilitation):
        return HttpResponseForbidden("No autorizado.")
    if not habilitation.certificate_pdf:
        raise Http404("Certificado no disponible.")
    filename = habilitation.certificate_pdf.name.rsplit("/", 1)[-1]
    return FileResponse(
        habilitation.certificate_pdf.open("rb"), as_attachment=True, filename=filename
    )


@login_required_icts
def lo3_habilitation_revoke(request, habilitation_id):
    if request.method != "POST":
        return HttpResponseForbidden("Metodo no permitido.")
    if not _is_lo3_principal(request.user):
        return HttpResponseForbidden("No autorizado.")

    habilitation = get_object_or_404(LO3Habilitation, pk=habilitation_id)
    reason = (request.POST.get("revoke_reason") or "").strip()
    habilitation.status = "revoked"
    habilitation.revoked_at = timezone.now()
    if reason:
        habilitation.revoke_reason = reason
    habilitation.save(update_fields=["status", "revoked_at", "revoke_reason"])
    messages.success(request, "Habilitacion revocada.")
    return redirect("sigmaconf:lo3_habilitation_detail", habilitation_id=habilitation.pk)


@login_required_icts
def lo3_section(request, section):
    is_tech, _ = _lo3_access_flags(request.user)
    if not is_tech:
        return HttpResponseForbidden("No autorizado.")

    sections = {
        "comunicaciones": {
            "title": "Comunicaciones con el cliente",
            "subtitle": "Seguimiento y registro de mensajes LO3.",
            "description": "Esta secci\u00f3n centralizar\u00e1 comunicaciones con solicitantes y entregas.",
            "notes": [
                "Historial de contactos vinculados a sesiones LO3.",
                "Registro de env\u00edos y respuestas del solicitante.",
            ],
            "nav_key": "communications",
        },
        "equipos": {
            "title": "Fichas de equipos",
            "subtitle": "Documentaci\u00f3n t\u00e9cnica y estado de instrumentos.",
            "description": "Aqu\u00ed se organizar\u00e1n manuales, fichas y control de equipos LO3.",
            "notes": [
                "Acceso r\u00e1pido a fichas t\u00e9cnicas.",
                "Historial de mantenimiento y verificaciones.",
            ],
            "nav_key": "equipment",
        },
        "calibraciones": {
            "title": "Calibraciones",
            "subtitle": "Interna y externa",
            "description": "Control de calibraciones LO3.",
            "notes": [
                "Calendario de calibraciones internas.",
                "Registro de certificados de calibraci\u00f3n externa.",
            ],
            "nav_key": "calibrations",
        },
        "materiales": {
            "title": "Materiales de referencia",
            "subtitle": "Registro de MR LO3.",
            "description": "Cat\u00e1logo de materiales de referencia LO3.",
            "notes": [
                "Listado de MR por lote.",
                "Trazabilidad de uso y reposici\u00f3n.",
            ],
            "nav_key": "materials",
        },
        "incidencias": {
            "title": "Controles intermedios / Incidencias",
            "subtitle": "Monitorizaci\u00f3n de incidencias y controles.",
            "description": "Panel para registrar incidencias operativas o controles intermedios.",
            "notes": [
                "Alertas de desviaciones en sesiones LO3.",
                "Registro de acciones correctivas.",
            ],
            "nav_key": "incidents",
        },
        "habilitaciones": {
            "title": "Habilitaciones",
            "subtitle": "Certificados y permisos del personal LO3.",
            "description": "Gestion de certificados de habilitacion LO3.",
            "notes": [
                "PDF de certificados de habilitaci\u00f3n.",
                "Historial de renovaciones.",
            ],
            "nav_key": "habilitations",
        },
    }

    config = sections.get(section)
    if not config:
        raise Http404("Seccion no encontrada.")

    context = {
        "section_title": config["title"],
        "section_subtitle": config["subtitle"],
        "section_description": config["description"],
        "section_notes": config.get("notes", []),
    }
    context.update(_lo3_nav_context(request.user, config["nav_key"]))
    return render(request, "sigmaconf/lo3_section.html", context)


@login_required_icts
def mcf_session_detail(request, pk):
    session = get_object_or_404(MCFSession, pk=pk)
    allowed, read_only = _can_view_session(request.user, session)
    if not allowed:
        messages.error(request, "No tienes permisos para acceder a esta sesion.")
        return redirect("sigmaconf:home")

    samples = list(session.samples.order_by("sequence"))
    attachments = (
        LO3Attachment.objects.filter(Q(session=session) | Q(sample__session=session))
        .select_related("sample", "uploaded_by")
        .order_by("-created_at")
    )
    total_samples = len(samples)
    extra_samples = sum(1 for sample in samples if sample.source == "session")
    declared_samples = total_samples - extra_samples
    analyzed_samples = sum(1 for sample in samples if sample.analysis_date)
    back_url = "sigmaconf:mcf_dashboard"
    if session.access_proposal.applicant_id == request.user.id:
        back_url = "icts:user_dashboard"
    elif read_only:
        back_url = "sigmaconf:home"
    context = {
        "session": session,
        "samples": samples,
        "attachments": attachments,
        "read_only": read_only,
        "back_url": back_url,
        "sample_stats": {
            "total": total_samples,
            "declared": declared_samples,
            "extra": extra_samples,
            "analyzed": analyzed_samples,
        },
    }
    if not read_only:
        context["final_report_form"] = MCFFinalReportUploadForm()
        context["attachment_form"] = LO3AttachmentForm()
    active_key = "results" if read_only else "requests"
    context.update(_lo3_nav_context(request.user, active_key))
    return render(request, "sigmaconf/mcf_session_detail.html", context)


@login_required_icts
def mcf_add_sample(request, pk):
    session = get_object_or_404(MCFSession, pk=pk)
    if not _can_edit_session(request.user, session):
        return HttpResponseForbidden("No autorizado.")

    if request.method == "POST":
        form = MCFSampleRecordForm(request.POST)
        if form.is_valid():
            sample = form.save(commit=False)
            sample.session = session
            sample.source = "session"
            max_seq = session.samples.aggregate(max_seq=Max("sequence")).get("max_seq") or 0
            sample.sequence = max_seq + 1
            sample.save()
            if sample.received_date and sample.analysis_date and sample.analysis_date < sample.received_date:
                messages.warning(request, "La fecha de analisis es anterior a la de recepcion.")
            messages.success(request, "Muestra anadida correctamente.")
            return redirect("sigmaconf:mcf_session_detail", pk=session.pk)
    else:
        form = MCFSampleRecordForm()

    return render(
        request,
        "sigmaconf/mcf_sample_edit.html",
        {"form": form, "session": session, "editing": False},
    )


@login_required_icts
def mcf_sample_edit(request, pk):
    sample = get_object_or_404(MCFSampleRecord, pk=pk)
    session = sample.session
    if not _can_edit_session(request.user, session):
        return HttpResponseForbidden("No autorizado.")

    if request.method == "POST":
        form = MCFSampleRecordForm(request.POST, instance=sample)
        if form.is_valid():
            sample = form.save()
            if sample.received_date and sample.analysis_date and sample.analysis_date < sample.received_date:
                messages.warning(request, "La fecha de analisis es anterior a la de recepcion.")
            messages.success(request, "Muestra actualizada.")
            return redirect("sigmaconf:mcf_session_detail", pk=session.pk)
    else:
        form = MCFSampleRecordForm(instance=sample)

    return render(
        request,
        "sigmaconf/mcf_sample_edit.html",
        {
            "form": form,
            "session": session,
            "editing": True,
            "sample": sample,
            "attachments": sample.attachments.select_related("uploaded_by"),
            "attachment_form": LO3AttachmentForm(),
        },
    )


@login_required_icts
def mcf_finish_session(request, pk):
    if request.method != "POST":
        return HttpResponseForbidden("Metodo no permitido.")

    session = get_object_or_404(MCFSession, pk=pk)
    was_completed = session.status == "completed"
    if not _can_edit_session(request.user, session):
        return HttpResponseForbidden("No autorizado.")

    samples = list(session.samples.order_by("sequence"))
    if not samples:
        messages.error(request, "No hay muestras registradas para finalizar la sesion.")
        return redirect("sigmaconf:mcf_session_detail", pk=session.pk)

    errors = []
    for sample in samples:
        label = f"Muestra {sample.sequence}"
        identification = (sample.identification or "").strip()
        if not identification:
            errors.append(f"{label}: falta la identificacion.")
        if not sample.analysis_date:
            errors.append(f"{label}: falta la fecha de analisis.")
        if not (sample.roughness or sample.image_2d or sample.image_3d or sample.thickness):
            errors.append(f"{label}: selecciona al menos una medida realizada.")

        if sample.analysis_date:
            if sample.return_date and sample.return_date < sample.analysis_date:
                errors.append(f"{label}: la fecha de devolucion es anterior al analisis.")
            if sample.report_delivery_date and sample.report_delivery_date < sample.analysis_date:
                errors.append(f"{label}: la entrega de informe es anterior al analisis.")

        sample_update_fields = []
        if sample.analysis_date and not sample.received_date:
            sample.received_date = sample.analysis_date
            sample_update_fields.append("received_date")
            note = "[AUTO] received_date asumida = analysis_date"
            observations = (sample.observations or "").strip()
            if note not in observations:
                sample.observations = f"{observations}\n{note}".strip() if observations else note
                sample_update_fields.append("observations")

        if sample_update_fields:
            sample.save(update_fields=sample_update_fields)

    if errors:
        for error in errors:
            messages.error(request, error)
        return redirect("sigmaconf:mcf_session_detail", pk=session.pk)

    update_fields = []
    if session.status != "completed":
        session.status = "completed"
        update_fields.append("status")
    if not session.completion_date:
        session.completion_date = timezone.now().date()
        update_fields.append("completion_date")

    regenerate = request.POST.get("regenerate") == "1"
    if not session.notice_pdf or regenerate:
        pdf_bytes = _build_notice_pdf(session)
        filename = f"Aviso_resultados_{session.lot_code}.pdf"
        session.notice_pdf.save(filename, ContentFile(pdf_bytes), save=False)
        update_fields.append("notice_pdf")

    if update_fields:
        session.save(update_fields=update_fields)

    if not was_completed and session.status == "completed":
        conversation, _ = LO3Conversation.objects.get_or_create(
            access_proposal=session.access_proposal
        )
        access_code = session.access_proposal.access_code or f"#{session.access_proposal_id}"
        results_url = reverse("sigmaconf:lo3_results")
        body = (
            "Resultados disponibles para la sesi\u00f3n "
            f"{session.lot_code} (solicitud {access_code}). "
            f"Puede consultarlos en Resultados: {results_url}"
        )
        LO3Message.objects.create(
            conversation=conversation,
            sender=None,
            body=body,
            is_system=True,
            read_at_technician=timezone.now(),
        )
        conversation.save(update_fields=["updated_at"])

    messages.success(request, "Sesion finalizada y aviso generado.")
    return redirect("sigmaconf:mcf_session_detail", pk=session.pk)


@login_required_icts
def mcf_notice_download(request, pk):
    session = get_object_or_404(MCFSession, pk=pk)
    allowed, _ = _can_view_session(request.user, session)
    if not allowed:
        return HttpResponseForbidden("No autorizado.")
    if not session.notice_pdf:
        raise Http404("Aviso no disponible.")
    filename = f"Aviso_resultados_{session.lot_code}.pdf"
    return FileResponse(session.notice_pdf.open("rb"), as_attachment=True, filename=filename)


@login_required_icts
def mcf_session_attachment_add(request, pk):
    if request.method != "POST":
        return HttpResponseForbidden("Metodo no permitido.")

    session = get_object_or_404(MCFSession, pk=pk)
    if not _can_edit_session(request.user, session):
        return HttpResponseForbidden("No autorizado.")

    form = LO3AttachmentForm(request.POST, request.FILES)
    if form.is_valid():
        attachment = form.save(commit=False)
        attachment.session = session
        attachment.uploaded_by = request.user
        attachment.save()
        messages.success(request, "Adjunto cargado.")
    else:
        messages.error(request, "No se pudo cargar el adjunto.")
    return redirect("sigmaconf:mcf_session_detail", pk=session.pk)


@login_required_icts
def mcf_sample_attachment_add(request, pk):
    if request.method != "POST":
        return HttpResponseForbidden("Metodo no permitido.")

    sample = get_object_or_404(MCFSampleRecord, pk=pk)
    session = sample.session
    if not _can_edit_session(request.user, session):
        return HttpResponseForbidden("No autorizado.")

    form = LO3AttachmentForm(request.POST, request.FILES)
    if form.is_valid():
        attachment = form.save(commit=False)
        attachment.sample = sample
        attachment.uploaded_by = request.user
        attachment.save()
        messages.success(request, "Adjunto cargado.")
    else:
        messages.error(request, "No se pudo cargar el adjunto.")
    return redirect("sigmaconf:mcf_sample_edit", pk=sample.pk)


@login_required_icts
def mcf_attachment_download(request, attachment_id):
    attachment = get_object_or_404(LO3Attachment, pk=attachment_id)
    session = attachment.session or attachment.sample.session
    allowed, _ = _can_view_session(request.user, session)
    if not allowed:
        return HttpResponseForbidden("No autorizado.")
    filename = attachment.file.name.rsplit("/", 1)[-1]
    return FileResponse(attachment.file.open("rb"), as_attachment=True, filename=filename)


@login_required_icts
def mcf_attachment_delete(request, attachment_id):
    if request.method != "POST":
        return HttpResponseForbidden("Metodo no permitido.")

    attachment = get_object_or_404(LO3Attachment, pk=attachment_id)
    session = attachment.session or attachment.sample.session
    if not _can_edit_session(request.user, session):
        return HttpResponseForbidden("No autorizado.")
    sample_id = attachment.sample_id
    attachment.delete()
    messages.success(request, "Adjunto eliminado.")
    if sample_id:
        return redirect("sigmaconf:mcf_sample_edit", pk=sample_id)
    return redirect("sigmaconf:mcf_session_detail", pk=session.pk)


@login_required_icts
def mcf_final_report_upload(request, pk):
    if request.method != "POST":
        return HttpResponseForbidden("Metodo no permitido.")

    session = get_object_or_404(MCFSession, pk=pk)
    if not _can_edit_session(request.user, session):
        return HttpResponseForbidden("No autorizado.")

    form = MCFFinalReportUploadForm(request.POST, request.FILES)
    if not form.is_valid():
        for error in form.errors.get("file", []):
            messages.error(request, error)
        return redirect("sigmaconf:mcf_session_detail", pk=session.pk)

    file_obj = form.cleaned_data["file"]
    filename = f"Informe_final_{session.lot_code}.pdf"
    session.final_report_pdf.save(filename, file_obj, save=False)
    session.final_report_uploaded_at = timezone.now()
    session.final_report_uploaded_by = request.user
    session.save(
        update_fields=[
            "final_report_pdf",
            "final_report_uploaded_at",
            "final_report_uploaded_by",
        ]
    )
    messages.success(request, "Informe final cargado correctamente.")
    return redirect("sigmaconf:mcf_session_detail", pk=session.pk)


@login_required_icts
def mcf_final_report_download(request, pk):
    session = get_object_or_404(MCFSession, pk=pk)
    allowed, _ = _can_view_session(request.user, session)
    if not allowed:
        return HttpResponseForbidden("No autorizado.")
    if not session.final_report_pdf:
        raise Http404("Informe final no disponible.")
    filename = f"Informe_final_{session.lot_code}.pdf"
    return FileResponse(
        session.final_report_pdf.open("rb"), as_attachment=True, filename=filename
    )


@login_required_icts
def mcf_export_excel(request, pk):
    session = get_object_or_404(MCFSession, pk=pk)
    allowed, _ = _can_view_session(request.user, session)
    if not allowed:
        return HttpResponseForbidden("No autorizado.")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "MCF"

    fallback_headers = [
        "Código solicitud",
        "Lote (MCF-aa-nnn)",
        "Identificación muestra",
        "Fecha recepción",
        "Fecha análisis",
        "Operador",
        "Rugosidad",
        "Imagen 2D",
        "Imagen 3D",
        "Espesor",
        "Fecha devolución",
        "Informe",
        "Fecha entrega informe",
        "Observaciones",
        "I1",
        "I2",
    ]
    mapping = load_quality_json(["lo3", "REV001", "mcf_register_mapping.json"], {})
    headers = mapping.get("headers") or fallback_headers
    sheet.append(headers)

    proposal = session.access_proposal
    access_code = proposal.access_code or f"#{proposal.id}"

    for sample in session.samples.order_by("sequence"):
        sheet.append([
            access_code,
            session.lot_code,
            sample.identification,
            sample.received_date,
            sample.analysis_date,
            sample.operator_name,
            "X" if sample.roughness else "",
            "X" if sample.image_2d else "",
            "X" if sample.image_3d else "",
            "X" if sample.thickness else "",
            sample.return_date,
            sample.report_code or session.report_code,
            sample.report_delivery_date,
            sample.observations,
            sample.indicator_i1,
            sample.indicator_i2,
        ])

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    filename = f"Registro_{session.lot_code}.xlsx"
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
