import re
import textwrap
from io import BytesIO

from django.contrib import messages
from django.conf import settings
from django.core.files.base import ContentFile
from django.db.models import Max, Prefetch
from django.http import FileResponse, Http404, HttpResponse, HttpResponseForbidden
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone

from icts.decorators import login_required_icts
from icts.auth_utils import (
    IMP_TECH_GROUPS,
    MANAGER_GROUPS,
    RESPONSABLE_GROUPS,
    get_normalized_user_groups,
    user_in_groups,
)
from icts.models import AccessProposal

from openpyxl import Workbook

from .forms import (
    IMPFinalReportUploadForm,
    IMPMessageForm,
    IMPSampleRecordForm,
    IMPSessionCreateForm,
)
from .models import (
    IMPConversation,
    IMPMessage,
    IMPMessageAttachment,
    IMPSampleRecord,
    IMPSession,
)


def _is_imp_technician(user):
    if not getattr(user, "is_authenticated", False):
        return False
    if user.is_staff or user.is_superuser:
        return True
    groups = get_normalized_user_groups(user)
    return user_in_groups(user, IMP_TECH_GROUPS, groups)


def _ensure_imp_technician(request):
    if not _is_imp_technician(request.user):
        messages.error(request, "No tienes permisos para acceder a esta seccion.")
        return False
    return True


def _is_imp_blocked(user, groups=None):
    if not getattr(user, "is_authenticated", False):
        return False
    if user.is_staff or user.is_superuser:
        return False
    groups = groups or get_normalized_user_groups(user)
    if user_in_groups(user, IMP_TECH_GROUPS, groups):
        return False
    return user_in_groups(user, RESPONSABLE_GROUPS | MANAGER_GROUPS, groups)


def _can_edit_session(user, session):
    groups = get_normalized_user_groups(user)
    if _is_imp_blocked(user, groups):
        return False
    is_tech = user_in_groups(user, IMP_TECH_GROUPS, groups)
    is_privileged = user.is_staff or user.is_superuser
    return (is_tech and session.technician_id == user.id) or is_privileged


def _can_view_session(user, session):
    groups = get_normalized_user_groups(user)
    if _is_imp_blocked(user, groups):
        return False, True
    is_tech = user_in_groups(user, IMP_TECH_GROUPS, groups)
    is_applicant = session.access_proposal.applicant_id == user.id
    is_privileged = user.is_staff or user.is_superuser
    allowed = is_applicant or is_tech or is_privileged
    can_edit = (is_tech and session.technician_id == user.id) or is_privileged
    read_only = not can_edit
    return allowed, read_only


def _can_access_imp_conversation(user, proposal):
    groups = get_normalized_user_groups(user)
    if _is_imp_blocked(user, groups):
        return False
    is_tech = user_in_groups(user, IMP_TECH_GROUPS, groups)
    is_privileged = user.is_staff or user.is_superuser
    is_applicant = proposal.applicant_id == user.id
    return is_applicant or is_tech or is_privileged


def _parse_imp_samples(imp_data):
    if not isinstance(imp_data, dict):
        return []
    pattern = re.compile(r"^imp_sample_(\d+)_(identification|name|details)$")
    grouped = {}
    for key, value in imp_data.items():
        match = pattern.match(key)
        if not match:
            continue
        idx = int(match.group(1))
        field = match.group(2)
        grouped.setdefault(idx, {})[field] = value
    return [grouped[idx] for idx in sorted(grouped.keys())]


def _build_notice_pdf(session):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buffer = BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4, pageCompression=0)
    height = A4[1]

    proposal = session.access_proposal
    access_code = proposal.access_code or f"#{proposal.id}"
    applicant_name = (
        proposal.contact_person
        or proposal.applicant.get_full_name()
        or proposal.applicant.get_username()
    )
    profile = getattr(proposal.applicant, "icts_profile", None)
    center = (getattr(profile, "center", "") or "").strip()
    address = (getattr(profile, "address", "") or "").strip()
    completion_date = (
        session.completed_at.date() if session.completed_at else timezone.now().date()
    )
    date_str = completion_date.strftime("%d/%m/%Y")

    default_title = "Responsable del servicio de Implantacion Ionica"
    default_name = "Dra. Teresa Hernandez Diaz"
    responsible_title = getattr(settings, "IMP_RESPONSIBLE_TITLE", default_title) or default_title
    responsible_name = getattr(settings, "IMP_RESPONSIBLE_NAME", default_name) or default_name

    def add_paragraph(text_obj, content):
        for line in textwrap.wrap(content, width=90):
            text_obj.textLine(line)

    text = pdf.beginText(50, height - 70)
    text.setFont("Helvetica", 11)
    text.textLine(f"A la atencion de D/Dna {applicant_name}")
    if center:
        text.textLine(center)
    if address:
        text.textLine(address)
    text.textLine("")
    text.textLine(f"Madrid, {date_str}")
    text.textLine("")
    text.textLine("Estimado Sr/Sra.:")
    text.textLine("")
    add_paragraph(
        text,
        "Adjunto le enviamos el resultado de las determinaciones efectuadas mediante "
        f"Implantacion Ionica del lote: {session.session_code}.",
    )
    text.textLine("")
    add_paragraph(
        text,
        "Puede encontrar dichos informes en su cuenta de usuario SIGMA ICTS, "
        f"asociados a la solicitud {access_code}.",
    )
    text.textLine("")
    text.textLine("Un cordial saludo")
    text.textLine("")
    text.textLine(responsible_title)
    text.textLine(responsible_name)

    pdf.drawText(text)
    pdf.showPage()
    pdf.save()
    buffer.seek(0)
    return buffer.read()


@login_required_icts
def imp_dashboard(request):
    groups = get_normalized_user_groups(request.user)
    if _is_imp_blocked(request.user, groups):
        return HttpResponseForbidden("No autorizado.")
    if not _ensure_imp_technician(request):
        return redirect("icts:dashboard")

    is_privileged = request.user.is_staff or request.user.is_superuser
    accepted_proposals = AccessProposal.objects.filter(
        status="accepted", facility_imp=True, imp_sessions__isnull=True
    ).order_by("-created_at")
    if is_privileged:
        all_sessions = IMPSession.objects.all()
    else:
        all_sessions = IMPSession.objects.filter(technician=request.user)
    sessions = all_sessions.filter(status="in_progress").order_by("-created_at")
    context = {
        "accepted_proposals": accepted_proposals,
        "sessions": sessions,
        "in_progress_count": sessions.count(),
        "completed_count": all_sessions.filter(status="completed").count(),
    }
    return render(request, "sigmaimp/imp_dashboard.html", context)


@login_required_icts
def imp_create_session(request, proposal_id=None):
    if not _ensure_imp_technician(request):
        return redirect("icts:dashboard")

    proposal = None
    if proposal_id:
        proposal = get_object_or_404(
            AccessProposal, pk=proposal_id, status="accepted", facility_imp=True
        )

    if request.method == "POST":
        if proposal is None:
            form = IMPSessionCreateForm(request.POST)
            if not form.is_valid():
                return render(
                    request,
                    "sigmaimp/imp_create_session.html",
                    {"form": form, "available_proposals": form.fields["access_proposal"].queryset},
                )
            proposal = form.cleaned_data["access_proposal"]

        session = IMPSession(access_proposal=proposal, technician=request.user)
        proposal_data = proposal.facility_data or {}
        if isinstance(proposal_data, dict):
            session.request_snapshot = proposal_data.get("imp", {}) or {}
        session.save()

        samples_data = _parse_imp_samples(session.request_snapshot)
        created = 0
        sequence = 1
        for sample in samples_data:
            identification = str(sample.get("identification") or "").strip()
            name = str(sample.get("name") or "").strip()
            details = str(sample.get("details") or "").strip()
            if not identification and not name:
                continue
            IMPSampleRecord.objects.create(
                session=session,
                sequence=sequence,
                source="proposal",
                identification=identification or name or f"Sample {sequence}",
                name=name,
                details=details,
            )
            sequence += 1
            created += 1

        if created == 0:
            IMPSampleRecord.objects.create(
                session=session,
                sequence=1,
                source="proposal",
                identification="Sample 1",
                name="",
                details="",
            )
            messages.warning(request, "Sesion creada con una muestra base sin datos importados.")
        else:
            messages.success(request, f"Sesion creada con {created} muestras importadas.")
        return redirect("sigmaimp:imp_session_detail", pk=session.pk)

    form = IMPSessionCreateForm(initial={"access_proposal": proposal} if proposal else None)
    if proposal:
        form.fields["access_proposal"].queryset = AccessProposal.objects.filter(pk=proposal.pk)

    available = AccessProposal.objects.filter(
        status="accepted", facility_imp=True
    ).order_by("-created_at")
    return render(
        request,
        "sigmaimp/imp_create_session.html",
        {"form": form, "available_proposals": available},
    )


@login_required_icts
def imp_results(request):
    groups = get_normalized_user_groups(request.user)
    if _is_imp_blocked(request.user, groups):
        return HttpResponseForbidden("No autorizado.")
    is_privileged = request.user.is_staff or request.user.is_superuser
    is_tech = user_in_groups(request.user, IMP_TECH_GROUPS, groups) or is_privileged

    sessions = IMPSession.objects.filter(status="completed")
    if is_privileged:
        pass
    elif is_tech:
        sessions = sessions.filter(technician=request.user)
    else:
        sessions = sessions.filter(access_proposal__applicant=request.user)

    sessions = sessions.select_related("access_proposal", "technician").order_by(
        "-completed_at", "-created_at"
    )
    context = {
        "sessions": sessions,
        "is_tech": is_tech,
        "is_privileged": is_privileged,
    }
    return render(request, "sigmaimp/imp_results.html", context)


@login_required_icts
def imp_notice_download(request, pk):
    session = get_object_or_404(IMPSession, pk=pk)
    allowed, _ = _can_view_session(request.user, session)
    if not allowed:
        return HttpResponseForbidden("No autorizado.")
    if not session.notice_pdf:
        raise Http404("Aviso no disponible.")
    filename = session.notice_pdf.name.rsplit("/", 1)[-1]
    return FileResponse(session.notice_pdf.open("rb"), as_attachment=True, filename=filename)


@login_required_icts
def imp_final_report_download(request, pk):
    session = get_object_or_404(IMPSession, pk=pk)
    allowed, _ = _can_view_session(request.user, session)
    if not allowed:
        return HttpResponseForbidden("No autorizado.")
    if not session.final_report_pdf:
        raise Http404("Informe final no disponible.")
    filename = session.final_report_pdf.name.rsplit("/", 1)[-1]
    return FileResponse(
        session.final_report_pdf.open("rb"), as_attachment=True, filename=filename
    )


@login_required_icts
def imp_export_excel(request, pk):
    session = get_object_or_404(IMPSession, pk=pk)
    allowed, _ = _can_view_session(request.user, session)
    if not allowed:
        return HttpResponseForbidden("No autorizado.")

    wb = Workbook()
    ws = wb.active
    ws.title = "Registro IMP"

    headers = [
        "Codigo solicitud",
        "Codigo sesion",
        "Identificacion muestra",
        "Nombre muestra",
        "Detalles",
        "Fecha recepcion",
        "Fecha implantacion",
        "Fecha devolucion",
        "Destruida",
        "Fecha destruccion",
        "Notas",
    ]
    ws.append(headers)

    access_code = session.access_proposal.access_code or f"#{session.access_proposal_id}"
    for sample in session.samples.order_by("sequence"):
        ws.append(
            [
                access_code,
                session.session_code,
                sample.identification,
                sample.name,
                sample.details,
                sample.received_date.strftime("%Y-%m-%d") if sample.received_date else "",
                sample.implant_date.strftime("%Y-%m-%d") if sample.implant_date else "",
                sample.return_date.strftime("%Y-%m-%d") if sample.return_date else "",
                "X" if sample.destroyed else "",
                sample.destroyed_at.strftime("%Y-%m-%d") if sample.destroyed_at else "",
                sample.notes,
            ]
        )

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    filename = f"IMP_{session.session_code}_registro.xlsx"
    response["Content-Disposition"] = f"attachment; filename={filename}"
    wb.save(response)
    return response


def _create_results_message(session):
    conversation, _ = IMPConversation.objects.get_or_create(
        access_proposal=session.access_proposal
    )
    access_code = session.access_proposal.access_code or f"#{session.access_proposal_id}"
    results_url = reverse("sigmaimp:imp_results")
    body = (
        "Resultados disponibles para la sesion "
        f"{session.session_code} (solicitud {access_code}). "
        f"Puede consultarlos en Resultados: {results_url}"
    )
    exists = IMPMessage.objects.filter(
        conversation=conversation,
        is_system=True,
        body=body,
    ).exists()
    if exists:
        return
    IMPMessage.objects.create(
        conversation=conversation,
        sender=None,
        body=body,
        is_system=True,
        read_at_technician=timezone.now(),
    )
    conversation.save(update_fields=["updated_at"])


@login_required_icts
def imp_finish_session(request, pk):
    if request.method != "POST":
        return HttpResponseForbidden("Metodo no permitido.")

    session = get_object_or_404(IMPSession, pk=pk)
    if not _can_edit_session(request.user, session):
        return HttpResponseForbidden("No autorizado.")

    samples = list(session.samples.order_by("sequence"))
    if not samples:
        messages.error(request, "No hay muestras registradas para finalizar la sesion.")
        return redirect("sigmaimp:imp_session_detail", pk=session.pk)

    errors = []
    for sample in samples:
        label = f"Muestra {sample.sequence}"
        identification = (sample.identification or "").strip()
        if not identification:
            errors.append(f"{label}: falta la identificacion.")
        if not sample.implant_date:
            errors.append(f"{label}: falta la fecha de implantacion.")
        if sample.received_date and sample.implant_date and sample.implant_date < sample.received_date:
            errors.append(f"{label}: implantacion anterior a la recepcion.")
        if sample.return_date and sample.implant_date and sample.return_date < sample.implant_date:
            errors.append(f"{label}: devolucion anterior a la implantacion.")
        if sample.destroyed and not sample.destroyed_at:
            errors.append(f"{label}: indica la fecha de destruccion.")
        if sample.destroyed_at and sample.implant_date and sample.destroyed_at < sample.implant_date:
            errors.append(f"{label}: destruccion anterior a la implantacion.")

    if errors:
        for error in errors:
            messages.error(request, error)
        return redirect("sigmaimp:imp_session_detail", pk=session.pk)

    was_completed = session.status == "completed"
    update_fields = []
    if session.status != "completed":
        session.status = "completed"
        update_fields.append("status")
    if not session.completed_at:
        session.completed_at = timezone.now()
        update_fields.append("completed_at")

    if not session.notice_pdf:
        pdf_bytes = _build_notice_pdf(session)
        filename = f"Aviso_resultados_{session.session_code}.pdf"
        session.notice_pdf.save(filename, ContentFile(pdf_bytes), save=False)
        update_fields.append("notice_pdf")

    if update_fields:
        session.save(update_fields=update_fields)

    if not was_completed and session.status == "completed":
        _create_results_message(session)

    messages.success(request, "Sesion finalizada y aviso generado.")
    return redirect("sigmaimp:imp_session_detail", pk=session.pk)


@login_required_icts
def imp_final_report_upload(request, pk):
    if request.method != "POST":
        return HttpResponseForbidden("Metodo no permitido.")

    session = get_object_or_404(IMPSession, pk=pk)
    if not _can_edit_session(request.user, session):
        return HttpResponseForbidden("No autorizado.")

    form = IMPFinalReportUploadForm(request.POST, request.FILES)
    if not form.is_valid():
        for error in form.errors.get("file", []):
            messages.error(request, error)
        return redirect("sigmaimp:imp_session_detail", pk=session.pk)

    file_obj = form.cleaned_data["file"]
    filename = f"Informe_final_{session.session_code}.pdf"
    session.final_report_pdf.save(filename, file_obj, save=False)
    session.final_report_uploaded_at = timezone.now()
    session.final_report_uploaded_by = request.user
    session.save(
        update_fields=[
            "final_report_pdf",
            "final_report_uploaded_at",
            "final_report_uploaded_by",
        ]
    )
    messages.success(request, "Informe final cargado correctamente.")
    return redirect("sigmaimp:imp_session_detail", pk=session.pk)


@login_required_icts
def imp_communications_list(request):
    groups = get_normalized_user_groups(request.user)
    if _is_imp_blocked(request.user, groups):
        return HttpResponseForbidden("No autorizado.")
    is_privileged = request.user.is_staff or request.user.is_superuser
    is_tech = user_in_groups(request.user, IMP_TECH_GROUPS, groups) or is_privileged
    is_staff_side = is_tech or is_privileged

    proposals = AccessProposal.objects.filter(
        status="accepted", facility_imp=True
    ).select_related("applicant")
    if not is_staff_side:
        proposals = proposals.filter(applicant=request.user)
    proposals = proposals.order_by("-created_at")

    conversations = (
        IMPConversation.objects.filter(access_proposal__in=proposals)
        .select_related("access_proposal", "access_proposal__applicant")
        .prefetch_related(
            Prefetch(
                "messages",
                queryset=IMPMessage.objects.select_related("sender").order_by("-created_at"),
            )
        )
    )
    conversation_map = {conversation.access_proposal_id: conversation for conversation in conversations}

    entries = []
    for proposal in proposals:
        conversation = conversation_map.get(proposal.id)
        last_message = None
        unread_count = 0
        if conversation:
            message_list = list(conversation.messages.all())
            if message_list:
                last_message = message_list[0]
                if is_staff_side:
                    unread_count = sum(
                        1 for msg in message_list if msg.read_at_technician is None
                    )
                else:
                    unread_count = sum(1 for msg in message_list if msg.read_at_applicant is None)
        entries.append(
            {
                "proposal": proposal,
                "conversation": conversation,
                "last_message": last_message,
                "unread_count": unread_count,
            }
        )

    context = {"entries": entries, "is_staff_side": is_staff_side}
    return render(request, "sigmaimp/imp_communications_list.html", context)


@login_required_icts
def imp_communications_detail(request, proposal_id):
    groups = get_normalized_user_groups(request.user)
    if _is_imp_blocked(request.user, groups):
        return HttpResponseForbidden("No autorizado.")
    proposal = get_object_or_404(
        AccessProposal, pk=proposal_id, status="accepted", facility_imp=True
    )
    if not _can_access_imp_conversation(request.user, proposal):
        return HttpResponseForbidden("No autorizado.")

    is_privileged = request.user.is_staff or request.user.is_superuser
    is_tech = user_in_groups(request.user, IMP_TECH_GROUPS, groups) or is_privileged
    is_staff_side = is_tech or is_privileged
    is_applicant = proposal.applicant_id == request.user.id

    conversation, _ = IMPConversation.objects.get_or_create(access_proposal=proposal)

    if request.method == "POST":
        form = IMPMessageForm(request.POST, request.FILES)
        if form.is_valid():
            body = (form.cleaned_data.get("body") or "").strip()
            if not body:
                messages.error(request, "El mensaje no puede estar vacio.")
            else:
                message = IMPMessage.objects.create(
                    conversation=conversation,
                    sender=request.user,
                    body=body,
                    read_at_applicant=timezone.now() if is_applicant else None,
                    read_at_technician=timezone.now() if is_staff_side else None,
                )
                for file_obj in request.FILES.getlist("attachments"):
                    IMPMessageAttachment.objects.create(
                        message=message,
                        file=file_obj,
                        filename=file_obj.name,
                        uploaded_by=request.user,
                    )
                conversation.save(update_fields=["updated_at"])
                messages.success(request, "Mensaje enviado.")
                return redirect(
                    "sigmaimp:imp_communications_detail", proposal_id=proposal.pk
                )
    else:
        form = IMPMessageForm()

    messages_qs = (
        conversation.messages.select_related("sender")
        .prefetch_related("attachments")
        .order_by("created_at")
    )

    if is_staff_side:
        unread_qs = messages_qs.filter(read_at_technician__isnull=True)
        if unread_qs.exists():
            unread_qs.update(read_at_technician=timezone.now())
    else:
        unread_qs = messages_qs.filter(read_at_applicant__isnull=True)
        if unread_qs.exists():
            unread_qs.update(read_at_applicant=timezone.now())

    context = {
        "proposal": proposal,
        "conversation": conversation,
        "messages": messages_qs,
        "form": form,
        "is_staff_side": is_staff_side,
    }
    return render(request, "sigmaimp/imp_communications_detail.html", context)


@login_required_icts
def imp_session_detail(request, pk):
    session = get_object_or_404(IMPSession, pk=pk)
    allowed, read_only = _can_view_session(request.user, session)
    if not allowed:
        messages.error(request, "No tienes permisos para acceder a esta sesion.")
        return redirect("icts:dashboard")

    samples = list(session.samples.order_by("sequence"))
    total_samples = len(samples)
    extra_samples = sum(1 for sample in samples if sample.source == "session")
    declared_samples = total_samples - extra_samples
    back_url = "sigmaimp:imp_dashboard"
    if session.access_proposal.applicant_id == request.user.id:
        back_url = "icts:user_dashboard"

    context = {
        "session": session,
        "samples": samples,
        "read_only": read_only,
        "back_url": back_url,
        "sample_stats": {
            "total": total_samples,
            "declared": declared_samples,
            "extra": extra_samples,
        },
    }
    if not read_only:
        context["final_report_form"] = IMPFinalReportUploadForm()
    return render(request, "sigmaimp/imp_session_detail.html", context)


@login_required_icts
def imp_add_sample(request, pk):
    session = get_object_or_404(IMPSession, pk=pk)
    if not _can_edit_session(request.user, session):
        return HttpResponseForbidden("No autorizado.")

    if request.method == "POST":
        form = IMPSampleRecordForm(request.POST)
        if form.is_valid():
            sample = form.save(commit=False)
            sample.session = session
            sample.source = "session"
            max_seq = session.samples.aggregate(max_seq=Max("sequence")).get("max_seq") or 0
            sample.sequence = max_seq + 1
            sample.save()
            messages.success(request, "Muestra anadida correctamente.")
            return redirect("sigmaimp:imp_session_detail", pk=session.pk)
    else:
        form = IMPSampleRecordForm()

    return render(
        request,
        "sigmaimp/imp_sample_edit.html",
        {"form": form, "session": session, "editing": False},
    )


@login_required_icts
def imp_sample_edit(request, pk):
    sample = get_object_or_404(IMPSampleRecord, pk=pk)
    session = sample.session
    if not _can_edit_session(request.user, session):
        return HttpResponseForbidden("No autorizado.")

    if request.method == "POST":
        form = IMPSampleRecordForm(request.POST, instance=sample)
        if form.is_valid():
            form.save()
            messages.success(request, "Muestra actualizada.")
            return redirect("sigmaimp:imp_session_detail", pk=session.pk)
    else:
        form = IMPSampleRecordForm(instance=sample)

    return render(
        request,
        "sigmaimp/imp_sample_edit.html",
        {"form": form, "session": session, "editing": True, "sample": sample},
    )
